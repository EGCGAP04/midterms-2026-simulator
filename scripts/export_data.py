#!/usr/bin/env python3
"""Export the forecast workbook to public/races.json.

Implements docs/SPEC.md §2: reads the cached values (and, in a second pass,
the formulas that tell the House race types apart), builds every race, runs
the self-checks of §2.6 and fails loudly if any of them does not hold.

Usage:
    python scripts/export_data.py <xlsx> --house-sheet "Senate House" \
        --senate-sheet "29-10 Senate" --out public/races.json
"""

from __future__ import annotations

import argparse
import json
import math
import re
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

# --- Workbook layout (docs/SPEC.md §2.2 and §2.4) ---------------------------

HOUSE_FIRST_ROW = 15
HOUSE_LAST_ROW = 883
HOUSE_RACES = 435
SENATE_FIRST_ROW = 7
SENATE_LAST_ROW = 41
SENATE_RACES = 35

# Bankhead is only named in Montana's note, not in any candidate cell.
MT3_THIRD_CANDIDATE = "Alani Bankhead"
MT3_NOTE_MUST_MENTION = "Bankhead"

STATES = {
    "Alabama": "AL", "Alaska": "AK", "Arizona": "AZ", "Arkansas": "AR",
    "California": "CA", "Colorado": "CO", "Connecticut": "CT", "Delaware": "DE",
    "Florida": "FL", "Georgia": "GA", "Hawaii": "HI", "Idaho": "ID",
    "Illinois": "IL", "Indiana": "IN", "Iowa": "IA", "Kansas": "KS",
    "Kentucky": "KY", "Louisiana": "LA", "Maine": "ME", "Maryland": "MD",
    "Massachusetts": "MA", "Michigan": "MI", "Minnesota": "MN",
    "Mississippi": "MS", "Missouri": "MO", "Montana": "MT", "Nebraska": "NE",
    "Nevada": "NV", "New Hampshire": "NH", "New Jersey": "NJ",
    "New Mexico": "NM", "New York": "NY", "North Carolina": "NC",
    "North Dakota": "ND", "Ohio": "OH", "Oklahoma": "OK", "Oregon": "OR",
    "Pennsylvania": "PA", "Rhode Island": "RI", "South Carolina": "SC",
    "South Dakota": "SD", "Tennessee": "TN", "Texas": "TX", "Utah": "UT",
    "Vermont": "VT", "Virginia": "VA", "Washington": "WA",
    "West Virginia": "WV", "Wisconsin": "WI", "Wyoming": "WY",
}

# Workbook label status → exported Incumbency (docs/SPEC.md §2.2).
INCUMBENCY = {
    "Incumbent": "incumbent",
    "Gained Incumbency": "gainedIncumbency",
    "Lost Incumbent": "lostIncumbent",
    "Non-Incumbent": "nonIncumbent",
}

RATING_COUNT_ROWS = {"Tossup": 8, "Lean": 9, "Likely": 10, "Safe": 11}


class ExportError(Exception):
    """A workbook value the exporter cannot interpret."""


# --- Pure helpers (unit-tested in scripts/test_export_data.py) --------------


def clean_name(raw: Any) -> tuple[str, bool] | None:
    """Return (display name, is independent), or None for "no candidate".

    Drops the footnote asterisks and the "(I)" marker: "Kevin Kiley* (I)**"
    becomes ("Kevin Kiley", True).
    """
    if raw is None:
        return None
    text = str(raw)
    independent = "(I)" in text
    name = " ".join(text.replace("(I)", " ").replace("*", " ").split())
    if name in ("", "None"):
        return None
    return name, independent


def parse_label(raw: Any) -> str:
    """Map a workbook label to its incumbency status, dropping footnote
    asterisks and the candidate-quality codes after a dash."""
    status = str(raw).split(" - ")[0].replace("*", "").strip().rstrip(".")
    if status not in INCUMBENCY:
        raise ExportError(f"unknown candidate label {raw!r}")
    return INCUMBENCY[status]


def district_id(name: str) -> tuple[str, str]:
    """Return (USPS state code, district ID) for a House district name:
    "California 22" → ("CA", "CA-22"); "Alaska at-large" and "Wyoming" →
    ("AK", "AK-AL") and ("WY", "WY-AL")."""
    match = re.fullmatch(r"(.+?) (\d+)", name)
    if match:
        state, number = match.group(1), match.group(2)
        code = STATES.get(state)
        if code is None:
            raise ExportError(f"unknown state in district {name!r}")
        return code, f"{code}-{int(number):02d}"
    state = name.removesuffix(" at-large")
    code = STATES.get(state)
    if code is None:
        raise ExportError(f"cannot parse district name {name!r}")
    return code, f"{code}-AL"


def senate_state(name: str) -> str:
    """USPS code for a Senate row: "Florida (special)" → "FL"."""
    state = name.removesuffix(" (special)")
    code = STATES.get(state)
    if code is None:
        raise ExportError(f"unknown state in Senate race {name!r}")
    return code


def excel_rating(margin: float) -> str:
    """The workbook's rating for a margin (R − opposition), e.g. "Lean D".
    Same thresholds as docs/SPEC.md §3.2; exactly 0 is Tilt R."""
    side = "R" if margin >= 0 else "D"
    size = abs(margin)
    if size >= 0.12:
        return f"Safe {side}"
    if size >= 0.06:
        return f"Likely {side}"
    if size >= 0.02:
        return f"Lean {side}"
    return f"Tossup (Tilt {side})"


def excel_flip(holder: str, margin: float) -> str:
    """The workbook's flip flag: whether the side ahead differs from the
    holder's side (an independent ahead counts as the D side)."""
    if margin >= 0:
        return "NO" if holder == "R" else "Yes, REP"
    return "Yes, DEM" if holder == "R" else "NO"


def blend_slope(at_row: float, at_companion: float) -> float:
    """Slope of a top-two/top-four primary blend: 1 / (1 + AT{r} / AT{r+1})."""
    return 1 / (1 + at_row / at_companion)


def mt3_shares(ac: float, ad: float, ae: float) -> tuple[float, float, float]:
    """Montana's baseline shares (A, B, C) for Alme, Bodnar and Bankhead."""
    m3 = (ac + ad) / 2
    a = (1 + m3) / 2
    c = ae - m3
    return a, 1 - a - c, c


def two_party_margin(rep_votes: float, dem_votes: float) -> float:
    """(R − D) / (R + D), with D votes given as a positive number."""
    return (rep_votes - dem_votes) / (rep_votes + dem_votes)


# --- Workbook reading --------------------------------------------------------


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def number(sheet: Any, cell: str) -> float:
    value = sheet[cell].value
    if not is_number(value):
        raise ExportError(f"{sheet.title}!{cell} should be a number, found {value!r}")
    return float(value)


def formula(sheet: Any, cell: str) -> str:
    value = sheet[cell].value
    return str(getattr(value, "text", value))


def candidate(name_raw: Any, label_raw: Any, default_party: str) -> dict[str, Any] | None:
    parsed = clean_name(name_raw)
    if parsed is None:
        return None
    name, independent = parsed
    return {"name": name, "label": parse_label(label_raw), "party": "I" if independent else default_party}


def read_house(values: Any, formulas: Any) -> tuple[list[dict[str, Any]], dict[str, list[float]]]:
    """House races, plus the 2024 presidential votes by state (Trump, Harris)."""
    factor = number(values, "AD2")  # PVI → margin factor of redrawn districts
    races = []
    votes: dict[str, list[float]] = {}
    for r in range(HOUSE_FIRST_ROW, HOUSE_LAST_ROW + 1, 2):
        name = str(values[f"C{r}"].value)
        state, race_id = district_id(name)
        m0 = number(values, f"AE{r}")
        aa_formula = formula(formulas, f"AA{r}")

        if "G$11" not in aa_formula:
            if clean_name(values[f"Y{r}"].value) is None:
                race_type = "uncontested"
            elif m0 == -1:
                race_type = "DvD"
            elif m0 == 1:
                race_type = "RvR"
            else:
                raise ExportError(f"{race_id}: fixed race with margin {m0}")
        elif re.search(r"\bAS\$?\d", aa_formula) and re.search(r"\bAT\$?\d", aa_formula):
            race_type = "blend"
        else:
            race_type = "normal"

        # In DvD / RvR races both slots belong to the same party.
        rep_party = {"DvD": "D"}.get(race_type, "R")
        opp_party = {"RvR": "R"}.get(race_type, "D")
        race: dict[str, Any] = {
            "id": race_id,
            "chamber": "house",
            "state": state,
            "name": name,
            "type": race_type,
            "holder": "R" if number(values, f"G{r + 1}") + number(values, f"J{r + 1}") >= 0 else "D",
            "rep": candidate(values[f"Y{r}"].value, values[f"Y{r + 1}"].value, rep_party),
            "opp": candidate(values[f"AB{r}"].value, values[f"AB{r + 1}"].value, opp_party),
            "m0": m0,
            "slope": blend_slope(number(values, f"AT{r}"), number(values, f"AT{r + 1}")) if race_type == "blend" else 1,
        }
        if race_type == "uncontested":
            race["fixedWinner"] = "R" if m0 > 0 else "D"
        elif race_type in ("DvD", "RvR"):
            race["fixedWinner"] = race_type[0]
            race["intraShares"] = [abs(number(values, f"AA{r}")), abs(number(values, f"AD{r}"))]

        # Context (docs/SPEC.md §2.2): the workbook puts the presidential result
        # in Q when there was no R-vs-D House race, so Q = S means "none".
        last_house, pres = number(values, f"Q{r}"), number(values, f"S{r}")
        new_pvi, old_pvi = values[f"U{r}"].value, values[f"V{r}"].value
        redrawn = is_number(old_pvi)
        year = 2022 if str(values[f"AU{r}"].value or "").strip() == "*2022" else 2024
        race["context"] = {
            "pres2024": pres + factor * (new_pvi - old_pvi) if redrawn else pres,
            "last": None if last_house == pres else {"margin": last_house, "year": year},
            "redrawn": redrawn and new_pvi != old_pvi,
        }

        # Every district partitions its state, so the votes add up statewide.
        state_votes = votes.setdefault(state, [0.0, 0.0])
        state_votes[0] += number(values, f"K{r + 1}")
        state_votes[1] += abs(number(values, f"N{r + 1}"))

        race["_check"] = {"rating": values[f"AF{r}"].value, "flip": values[f"AH{r}"].value}
        races.append(race)
    return races, votes


def read_senate(values: Any, formulas: Any, pres_votes: dict[str, list[float]]) -> list[dict[str, Any]]:
    races = []
    for r in range(SENATE_FIRST_ROW, SENATE_LAST_ROW + 1):
        name = str(values[f"B{r}"].value)
        state = senate_state(name)
        m0 = number(values, f"AE{r}")
        three_way = formula(formulas, f"AE{r}").replace(" ", "") != f"=AVERAGE(AC{r},AD{r})"
        race: dict[str, Any] = {
            "id": f"{state}-SEN",
            "chamber": "senate",
            "state": state,
            "name": name,
            "type": "MT3" if three_way else "normal",
            "holder": str(values[f"J{r}"].value),
            "rep": candidate(values[f"K{r}"].value, values[f"L{r}"].value, "R"),
            "opp": candidate(values[f"M{r}"].value, values[f"N{r}"].value, "D"),
            "m0": m0,
            "slope": 1,
        }
        if three_way:
            note = str(values[f"AH{r}"].value or "")
            if state != "MT" or MT3_NOTE_MUST_MENTION not in note:
                raise ExportError(f"{name}: unexpected three-way race formula in AE{r}")
            a, b, c = mt3_shares(number(values, f"AC{r}"), number(values, f"AD{r}"), m0)
            race["shares"] = {
                "A": a, "B": b, "C": c,
                "names": [race["rep"]["name"], race["opp"]["name"], MT3_THIRD_CANDIDATE],
            }
        poll_source = values[f"U{r}"].value
        if poll_source:
            race["pollSource"] = str(poll_source)
        trump, harris = pres_votes[state]
        race["context"] = {
            "pres2024": two_party_margin(trump, harris),
            "last": {"margin": number(values, f"C{r}"), "year": int(number(values, f"D{r}"))},
        }
        race["_check"] = {"rating": values[f"AF{r}"].value, "flip": values[f"AG{r}"].value, "row": r}
        races.append(race)
    return races


# --- Self-checks (docs/SPEC.md §2.6) ----------------------------------------


def baseline_winner(race: dict[str, Any]) -> str:
    """Party ahead at the baseline (fixed races: their fixed winner)."""
    if "fixedWinner" in race:
        return race["fixedWinner"]
    if race["type"] == "MT3":
        shares = race["shares"]
        best = max(("A", "B", "C"), key=lambda k: shares[k])
        return {"A": race["rep"]["party"], "B": race["opp"]["party"], "C": "D"}[best]
    return race["rep"]["party"] if race["m0"] >= 0 else race["opp"]["party"]


def self_check(house: list[dict[str, Any]], senate: list[dict[str, Any]], hv: Any, sv: Any, sf: Any, house_sheet: str) -> list[str]:
    errors = []
    if len(house) != HOUSE_RACES:
        errors.append(f"expected {HOUSE_RACES} House races, found {len(house)}")
    if len(senate) != SENATE_RACES:
        errors.append(f"expected {SENATE_RACES} Senate races, found {len(senate)}")
    ids = [race["id"] for race in house + senate]
    duplicates = sorted({i for i in ids if ids.count(i) > 1})
    if duplicates:
        errors.append(f"duplicate race IDs: {duplicates}")

    # The Senate sheet must reference the House sheet given on the command
    # line, so both chambers share the same baseline (§2.1).
    reference = formula(sf, f"V{SENATE_FIRST_ROW}")
    if f"'{house_sheet}'!" not in reference:
        errors.append(f"the Senate sheet does not reference '{house_sheet}' (V{SENATE_FIRST_ROW} is {reference})")

    for race in house + senate:
        expected = race["_check"]
        if excel_rating(race["m0"]) != expected["rating"]:
            errors.append(f"{race['id']}: rating {excel_rating(race['m0'])!r} != workbook {expected['rating']!r}")
        if excel_flip(race["holder"], race["m0"]) != expected["flip"]:
            errors.append(f"{race['id']}: flip {excel_flip(race['holder'], race['m0'])!r} != workbook {expected['flip']!r}")
        sides = [race[side] for side in ("rep", "opp") if race[side] is not None]
        if len(sides) != (1 if race["type"] == "uncontested" else 2):
            errors.append(f"{race['id']}: {len(sides)} candidate(s) in a {race['type']} race")

    # House totals and rating counts: R column = D, S column = R.
    house_d = sum(1 for race in house if race["m0"] < 0)
    if (house_d, len(house) - house_d) != (number(hv, "R7"), number(hv, "S7")):
        errors.append(f"House D/R {house_d}/{len(house) - house_d} != workbook {hv['R7'].value}/{hv['S7'].value}")
    counts: dict[tuple[str, str], int] = {}
    for race in house:
        key = (excel_rating(race["m0"]).split()[0], "R" if race["m0"] >= 0 else "D")
        counts[key] = counts.get(key, 0) + 1
    for bucket, row in RATING_COUNT_ROWS.items():
        for column, side in (("R", "D"), ("S", "R")):
            count = counts.get((bucket, side), 0)
            if count != number(hv, f"{column}{row}"):
                errors.append(f"House {bucket} {side}: {count} != workbook {column}{row} = {hv[f'{column}{row}'].value}")

    # Senate totals, including the seats not up.
    totals = {"D": number(sv, "N3"), "I": 0.0, "R": number(sv, "P3")}
    for race in senate:
        totals[baseline_winner(race)] += 1
    expected_totals = {"D": number(sv, "J44"), "I": number(sv, "L44"), "R": number(sv, "N44")}
    if totals != expected_totals:
        errors.append(f"Senate totals {totals} != workbook J44/L44/N44 {expected_totals}")

    for race in senate:
        if race["type"] == "MT3":
            shares = race["shares"]
            if not math.isclose(shares["A"] - shares["B"], race["m0"], abs_tol=1e-9):
                errors.append(f"{race['id']}: A − B = {shares['A'] - shares['B']} != AE{race['_check']['row']} = {race['m0']}")
    if sum(1 for race in senate if race["type"] == "MT3") != 1:
        errors.append("expected exactly one three-way (MT3) Senate race")
    return errors


# --- Output -------------------------------------------------------------------


def write_dataset(path: Path, dataset: dict[str, Any]) -> None:
    """Write the JSON with one race per line, so a new forecast shows up as a
    readable diff."""
    header = {key: value for key, value in dataset.items() if key != "races"}
    lines = ["{"]
    for key, value in header.items():
        lines.append(f"  {json.dumps(key)}: {json.dumps(value, ensure_ascii=False)},")
    lines.append('  "races": [')
    races = [json.dumps(race, ensure_ascii=False) for race in dataset["races"]]
    lines.append(",\n".join(f"    {race}" for race in races))
    lines.append("  ]")
    lines.append("}")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Export the forecast workbook to races.json (docs/SPEC.md §2).")
    parser.add_argument("xlsx", type=Path, help="the forecast workbook")
    parser.add_argument("--house-sheet", required=True, help='House sheet, e.g. "Senate House"')
    parser.add_argument("--senate-sheet", required=True, help='Senate sheet, e.g. "29-10 Senate"')
    parser.add_argument("--out", type=Path, default=Path("public/races.json"), help="output file (default: %(default)s)")
    args = parser.parse_args(argv)

    import openpyxl  # imported here so the pure helpers can be tested without it

    values_book = openpyxl.load_workbook(args.xlsx, data_only=True)
    formulas_book = openpyxl.load_workbook(args.xlsx, data_only=False)
    for sheet in (args.house_sheet, args.senate_sheet):
        if sheet not in values_book.sheetnames:
            parser.error(f"no sheet named {sheet!r}; the workbook has {values_book.sheetnames}")
    hv, hf = values_book[args.house_sheet], formulas_book[args.house_sheet]
    sv, sf = values_book[args.senate_sheet], formulas_book[args.senate_sheet]

    try:
        house, pres_votes = read_house(hv, hf)
        senate = read_senate(sv, sf, pres_votes)
        errors = self_check(house, senate, hv, sv, sf, args.house_sheet)
    except ExportError as error:
        print(f"error: {error}", file=sys.stderr)
        return 1

    if errors:
        print(f"error: {len(errors)} self-check(s) failed:", file=sys.stderr)
        for error in errors:
            print(f"  - {error}", file=sys.stderr)
        return 1

    for race in house + senate:
        del race["_check"]
    dataset = {
        "generatedAt": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "sourceFile": args.xlsx.name,
        "g0": number(hv, "G11"),
        "senateNotUp": {"D": int(number(sv, "N3")), "R": int(number(sv, "P3"))},
        "races": house + senate,
    }
    write_dataset(args.out, dataset)
    print(f"wrote {args.out}: {len(house)} House and {len(senate)} Senate races, all self-checks passed")
    return 0


if __name__ == "__main__":
    sys.exit(main())
